#!/usr/bin/env python3
"""
Home network inventory -> Excel.

Finds devices with several methods, because no single one sees everything:
  * arp-scan   Layer-2 ARP sweep. Every device on the LAN must answer ARP, even
               ones that ignore ping (most smart plugs, cameras, TVs, phones).
  * nmap -sn   ARP + ICMP + TCP/UDP probes, plus reverse-DNS names.
  * ip neigh   Devices this host has talked to recently (kernel ARP cache).
  * mDNS       avahi-browse: .local names, services (Chromecast, AirPlay,
               HomeKit, printers...) and friendly names / models.
  * SSDP/UPnP  Smart TVs, routers, media servers announce make/model here.
  * DNS        Reverse lookup against your DNS server (router / Pi-hole names).
  * NetBIOS    Windows / Samba computer names.
Optional: --ports (top TCP ports) and --os (nmap OS guess, slow).

Results are merged per MAC address and written to an .xlsx. A JSON history
file next to it keeps first-seen / last-seen dates, and the "Friendly name"
and "Notes" columns you type into the spreadsheet are kept on the next run.

Must run as root (raw sockets for ARP).
"""
import argparse
import concurrent.futures as cf
import datetime as dt
import ipaddress
import json
import os
import re
import shutil
import socket
import subprocess
import sys
import urllib.request
import xml.etree.ElementTree as ET
from collections import Counter

NOW = dt.datetime.now().replace(microsecond=0)


# ----------------------------------------------------------------- helpers --
def run(cmd, timeout=180):
    """Run a command, return stdout ('' if the tool is missing or fails)."""
    if not shutil.which(cmd[0]):
        log(f"  - {cmd[0]} not installed, skipping")
        return ""
    try:
        return subprocess.run(cmd, capture_output=True, text=True, timeout=timeout).stdout
    except (subprocess.TimeoutExpired, OSError) as e:
        log(f"  - {cmd[0]} failed: {e}")
        return ""


def log(msg):
    print(msg, file=sys.stderr, flush=True)


def norm_mac(mac):
    mac = (mac or "").strip().lower().replace("-", ":")
    return mac if re.fullmatch(r"([0-9a-f]{2}:){5}[0-9a-f]{2}", mac) else ""


def is_private_mac(mac):
    """Locally-administered bit set = randomised MAC (phones, laptops, tablets)."""
    return bool(mac) and int(mac.split(":")[0], 16) & 0x02 != 0


def ip_key(ip):
    try:
        return int(ipaddress.ip_address(ip))
    except ValueError:
        return 0


def default_iface():
    out = run(["ip", "-4", "route", "show", "default"])
    m = re.search(r"\bdev (\S+)", out)
    return m.group(1) if m else ""


def iface_info(iface):
    """(ip, mac) of the scanning host itself - arp-scan never lists itself."""
    out = run(["ip", "-o", "-4", "addr", "show", "dev", iface])
    ip = re.search(r"inet (\d+\.\d+\.\d+\.\d+)", out)
    link = run(["ip", "-o", "link", "show", "dev", iface])
    mac = re.search(r"link/ether (\S+)", link)
    return (ip.group(1) if ip else ""), norm_mac(mac.group(1) if mac else "")


# ----------------------------------------------------------------- sources --
def src_arp_scan(iface, subnet, retries):
    log(f"* arp-scan {subnet} on {iface} ({retries} retries)")
    out = run(["arp-scan", f"--interface={iface}", f"--retry={retries}", "--plain", subnet], timeout=300)
    found = {}
    for line in out.splitlines():
        parts = line.split("\t")
        if len(parts) >= 2 and re.fullmatch(r"\d+\.\d+\.\d+\.\d+", parts[0]):
            vendor = re.sub(r"\s*\(DUP: \d+\)$", "", parts[2]).strip() if len(parts) > 2 else ""
            found.setdefault(parts[0], {"mac": norm_mac(parts[1]), "vendor": vendor})
    log(f"  {len(found)} devices")
    return found


def src_nmap_ping(subnet, dns):
    log(f"* nmap host discovery {subnet}")
    cmd = ["nmap", "-sn", "-PR", "-PE", "-PP", "-PS21,22,23,53,80,443,554,1883,8008,8080,8443,9100",
           "-PA80,443", "-PU53,137,161,1900,5353", "-T4", "-oX", "-"]
    if dns:
        cmd += ["--dns-servers", dns]
    out = run(cmd + [subnet], timeout=600)
    found = {}
    if not out.strip():
        return found
    for h in ET.fromstring(out).iter("host"):
        if h.find("status") is not None and h.find("status").get("state") != "up":
            continue
        ip = mac = vendor = name = ""
        for a in h.findall("address"):
            if a.get("addrtype") == "ipv4":
                ip = a.get("addr")
            elif a.get("addrtype") == "mac":
                mac, vendor = norm_mac(a.get("addr")), a.get("vendor", "")
        hn = h.find("hostnames/hostname")
        if hn is not None:
            name = hn.get("name", "")
        if ip:
            found[ip] = {"mac": mac, "vendor": vendor, "dns": name}
    log(f"  {len(found)} devices")
    return found


def src_ip_neigh(subnet):
    net = ipaddress.ip_network(subnet, strict=False)
    found = {}
    for line in run(["ip", "-4", "neigh", "show"]).splitlines():
        m = re.match(r"(\d+\.\d+\.\d+\.\d+) .*lladdr (\S+) (\S+)", line)
        if m and ipaddress.ip_address(m.group(1)) in net and m.group(3) != "FAILED":
            found[m.group(1)] = {"mac": norm_mac(m.group(2))}
    log(f"* ip neigh: {len(found)} cached neighbours")
    return found


MDNS_TYPES = {
    "_googlecast._tcp": "Chromecast / Google Cast",
    "_airplay._tcp": "AirPlay",
    "_raop._tcp": "AirPlay speaker",
    "_hap._tcp": "HomeKit accessory",
    "_homekit._tcp": "HomeKit",
    "_matter._tcp": "Matter device",
    "_matterc._udp": "Matter device",
    "_ipp._tcp": "Printer",
    "_ipps._tcp": "Printer",
    "_printer._tcp": "Printer",
    "_pdl-datastream._tcp": "Printer",
    "_scanner._tcp": "Scanner",
    "_uscan._tcp": "Scanner",
    "_spotify-connect._tcp": "Spotify Connect speaker",
    "_sonos._tcp": "Sonos speaker",
    "_amzn-wplay._tcp": "Amazon Fire TV / Echo",
    "_androidtvremote2._tcp": "Android TV",
    "_roku._tcp": "Roku",
    "_hue._tcp": "Philips Hue bridge",
    "_esphomelib._tcp": "ESPHome device",
    "_home-assistant._tcp": "Home Assistant",
    "_smb._tcp": "File share (SMB)",
    "_afpovertcp._tcp": "Mac file share",
    "_device-info._tcp": "",
    "_companion-link._tcp": "Apple device",
    "_sleep-proxy._udp": "Apple TV / HomePod",
    "_workstation._tcp": "Computer",
    "_ssh._tcp": "SSH",
    "_http._tcp": "Web UI",
}


def src_mdns():
    log("* mDNS (avahi-browse)")
    out = run(["avahi-browse", "-a", "-r", "-t", "-p"], timeout=30)
    found = {}
    for line in out.splitlines():
        if not line.startswith("="):
            continue
        f = line.split(";")
        if len(f) < 9 or f[2] != "IPv4":
            continue
        name, stype, host, addr, txt = f[3], f[4], f[6], f[7], ";".join(f[9:])
        d = found.setdefault(addr, {"mdns": "", "services": set(), "model": "", "friendly": ""})
        d["mdns"] = d["mdns"] or host
        label = MDNS_TYPES.get(stype, stype.strip("_").split("._")[0])
        if label:
            d["services"].add(label)
        # Google Cast / HomeKit / AirPlay put useful details in TXT records
        fn = re.search(r'"fn=([^"]+)"', txt)
        md = re.search(r'"(?:md|model|am)=([^"]+)"', txt)
        if fn:
            d["friendly"] = d["friendly"] or fn.group(1)
        if md:
            d["model"] = d["model"] or md.group(1)
        if not d["friendly"] and stype in ("_googlecast._tcp", "_airplay._tcp", "_hap._tcp", "_raop._tcp"):
            d["friendly"] = re.sub(r"^[0-9A-F]{12}@", "", name.replace("\\032", " "))
    log(f"  {len(found)} devices advertise services")
    return found


def src_ssdp(wait=4):
    log("* SSDP / UPnP discovery")
    msg = ("M-SEARCH * HTTP/1.1\r\nHOST: 239.255.255.250:1900\r\n"
           'MAN: "ssdp:discover"\r\nMX: 2\r\nST: ssdp:all\r\n\r\n').encode()
    locations = {}
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM, socket.IPPROTO_UDP)
        s.setsockopt(socket.IPPROTO_IP, socket.IP_MULTICAST_TTL, 2)
        s.settimeout(wait)
        for _ in range(2):
            s.sendto(msg, ("239.255.255.250", 1900))
        end = dt.datetime.now() + dt.timedelta(seconds=wait)
        while dt.datetime.now() < end:
            try:
                data, (ip, _) = s.recvfrom(65507)
            except socket.timeout:
                break
            m = re.search(rb"(?im)^location:\s*(\S+)", data)
            if m:
                locations.setdefault(ip, set()).add(m.group(1).decode(errors="ignore"))
        s.close()
    except OSError as e:
        log(f"  - SSDP failed: {e}")
        return {}

    def describe(ip, urls):
        for url in sorted(urls):
            try:
                with urllib.request.urlopen(url, timeout=3) as r:
                    root = ET.fromstring(r.read())
                ns = {"u": root.tag.split("}")[0].strip("{")} if root.tag.startswith("{") else {}
                dev = root.find("u:device", ns) if ns else root.find("device")
                if dev is None:
                    continue
                g = (lambda t: (dev.findtext(f"u:{t}", "", ns) if ns else dev.findtext(t, "")).strip())
                return ip, {"upnp_name": g("friendlyName"),
                            "upnp_model": " ".join(x for x in (g("manufacturer"), g("modelName")) if x)}
            except Exception:
                continue
        return ip, {}

    found = {}
    with cf.ThreadPoolExecutor(16) as ex:
        for ip, info in ex.map(lambda kv: describe(*kv), locations.items()):
            if info:
                found[ip] = info
    log(f"  {len(locations)} UPnP responders, {len(found)} described")
    return found


def src_reverse_dns(ips, dns):
    log(f"* reverse DNS for {len(ips)} addresses" + (f" via {dns}" if dns else ""))

    def lookup(ip):
        if dns and shutil.which("dig"):
            out = run(["dig", "+short", "+time=2", "+tries=1", "-x", ip, f"@{dns}"], timeout=10).strip()
            name = out.splitlines()[0].rstrip(".") if out else ""
        else:
            try:
                name = socket.gethostbyaddr(ip)[0]
            except OSError:
                name = ""
        return ip, name

    with cf.ThreadPoolExecutor(32) as ex:
        return {ip: n for ip, n in ex.map(lookup, ips) if n}


def src_netbios(ips):
    if not shutil.which("nmblookup"):
        return {}
    log(f"* NetBIOS names for {len(ips)} addresses")

    def lookup(ip):
        out = run(["nmblookup", "-A", ip], timeout=6)
        m = re.search(r"^\s+(\S+)\s+<00> -\s+[BMH] <ACTIVE>", out, re.M)
        return ip, (m.group(1) if m else "")

    with cf.ThreadPoolExecutor(32) as ex:
        return {ip: n for ip, n in ex.map(lookup, ips) if n}


def src_nmap_detail(ips, ports, os_guess):
    if not ips or not (ports or os_guess):
        return {}
    log(f"* nmap {'ports ' if ports else ''}{'OS ' if os_guess else ''}scan of {len(ips)} hosts")
    cmd = ["nmap", "-Pn", "-T4", "-oX", "-"]
    cmd += ["-sS", "--top-ports", "100", "--open"] if ports else ["-sn"]
    if os_guess:
        cmd += ["-O", "--osscan-guess", "--max-os-tries", "1"]
    out = run(cmd + sorted(ips, key=ip_key), timeout=3600)
    found = {}
    if not out.strip():
        return found
    for h in ET.fromstring(out).iter("host"):
        ip = next((a.get("addr") for a in h.findall("address") if a.get("addrtype") == "ipv4"), "")
        pl = [f"{p.get('portid')}/{(p.find('service').get('name') if p.find('service') is not None else '?')}"
              for p in h.findall("ports/port") if p.find("state").get("state") == "open"]
        osm = h.find("os/osmatch")
        found[ip] = {"ports": ", ".join(pl), "os": f"{osm.get('name')} ({osm.get('accuracy')}%)" if osm is not None else ""}
    return found


# ------------------------------------------------------------ device guess --
VENDOR_HINTS = [
    (r"proxmox", "Proxmox VM"),
    (r"espressif", "IoT (ESP32/ESP8266 smart device)"),
    (r"tuya", "Smart plug / bulb (Tuya)"),
    (r"amazon", "Amazon Echo / Fire / Ring"),
    (r"google", "Google / Nest device"),
    (r"sonos", "Sonos speaker"),
    (r"roku", "Roku"),
    (r"philips|signify", "Philips Hue"),
    (r"ring", "Ring camera / doorbell"),
    (r"wyze", "Wyze camera"),
    (r"ecobee", "Ecobee thermostat"),
    (r"nest", "Nest device"),
    (r"samsung", "Samsung (TV / phone / appliance)"),
    (r"\blg\b|lg electronics", "LG (TV / appliance)"),
    (r"sony", "Sony (TV / PlayStation)"),
    (r"nintendo", "Nintendo console"),
    (r"microsoft", "Microsoft (Xbox / Surface)"),
    (r"apple", "Apple device"),
    (r"raspberry", "Raspberry Pi"),
    (r"hewlett|hp inc|epson|brother|canon", "Printer"),
    (r"ubiquiti|tp-link|netgear|asus|linksys|eero|cisco|mikrotik|d-link", "Network equipment"),
    (r"intel|realtek|dell|lenovo|asrock|gigabyte|micro-star", "Computer"),
    (r"lorex|hikvision|dahua|reolink|amcrest|axis", "Camera / NVR"),
]


def guess_type(d):
    for s in ("Chromecast / Google Cast", "Android TV", "Amazon Fire TV / Echo", "Roku", "Sonos speaker",
              "Printer", "Philips Hue bridge", "HomeKit accessory", "Matter device", "ESPHome device",
              "Home Assistant", "Apple TV / HomePod", "AirPlay speaker", "Spotify Connect speaker"):
        if s in d["services"]:
            return s
    if d["mac"].startswith("bc:24:11"):          # Proxmox's own OUI for VM NICs
        return "Proxmox VM"
    v = d["vendor"].lower()
    for pat, label in VENDOR_HINTS:
        if re.search(pat, v):
            return label
    if is_private_mac(d["mac"]):
        return "Phone / tablet / laptop (private MAC)"
    if d["upnp_model"]:
        return "UPnP device"
    return ""


# ---------------------------------------------------------------- history --
COLUMNS = [
    ("Status", 9), ("IP address", 15), ("Hostname", 28), ("MAC address", 19), ("Vendor", 26),
    ("Device type (guess)", 30), ("Friendly name (yours)", 24), ("Notes (yours)", 30),
    ("mDNS name", 26), ("DNS name", 28), ("NetBIOS name", 16), ("UPnP / model", 30),
    ("Advertised services", 34), ("Open ports", 30), ("OS guess", 28),
    ("Private MAC", 11), ("Found by", 26), ("First seen", 19), ("Last seen", 19),
]


def load_previous(json_path, xlsx_path):
    hist = {}
    if os.path.exists(json_path):
        with open(json_path) as f:
            hist = json.load(f)
    # Pick up anything you typed into the spreadsheet since the last run
    if os.path.exists(xlsx_path):
        try:
            from openpyxl import load_workbook
            ws = load_workbook(xlsx_path)["Devices"]
            hdr = [c.value for c in ws[1]]
            ci = {h: i for i, h in enumerate(hdr)}
            for row in ws.iter_rows(min_row=2, values_only=True):
                mac, ip = row[ci["MAC address"]] or "", row[ci["IP address"]]
                key = next((k for k in (f"{mac}|{ip}", mac, f"ip:{ip}") if k in hist and k != "|"), None)
                if key is None:
                    continue
                hist[key]["friendly_user"] = row[ci["Friendly name (yours)"]] or ""
                hist[key]["notes"] = row[ci["Notes (yours)"]] or ""
        except Exception as e:
            log(f"  - could not read your notes from {xlsx_path}: {e}")
    return hist


# ------------------------------------------------------------------- main --
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--subnet", required=True, help="e.g. 192.168.4.0/23")
    ap.add_argument("--iface", default="", help="network interface (default: the one with the default route)")
    ap.add_argument("--dns", default="", help="DNS server for reverse lookups (router / Pi-hole)")
    ap.add_argument("--out", default="reports/network_inventory.xlsx")
    ap.add_argument("--retries", type=int, default=3, help="arp-scan retries (more = catches sleepy devices)")
    ap.add_argument("--ports", action="store_true", help="also scan the top 100 TCP ports (slower)")
    ap.add_argument("--os", action="store_true", help="also guess the OS with nmap (much slower)")
    ap.add_argument("--keep-days", type=int, default=90, help="drop devices not seen for this many days")
    a = ap.parse_args()

    if os.geteuid() != 0:
        sys.exit("Run as root (ARP scanning needs raw sockets): sudo " + " ".join(sys.argv))

    iface = a.iface or default_iface()
    self_ip, self_mac = iface_info(iface)
    log(f"Scanning {a.subnet} from {self_ip or '?'} ({iface})")

    sources = {}
    sources["arp-scan"] = src_arp_scan(iface, a.subnet, a.retries)
    sources["nmap"] = src_nmap_ping(a.subnet, a.dns)
    sources["arp cache"] = src_ip_neigh(a.subnet)
    sources["mDNS"] = src_mdns()
    sources["UPnP"] = src_ssdp()

    net = ipaddress.ip_network(a.subnet, strict=False)
    ips = {ip for s in sources.values() for ip in s if ipaddress.ip_address(ip) in net}
    if self_ip:
        ips.add(self_ip)
    sources["DNS"] = src_reverse_dns(ips, a.dns)
    sources["NetBIOS"] = src_netbios(ips)
    detail = src_nmap_detail(ips, a.ports, a.os)

    # ---- merge everything per IP
    devices = {}
    for ip in ips:
        d = {"ip": ip, "mac": "", "vendor": "", "dns": "", "mdns": "", "netbios": "", "upnp_name": "",
             "upnp_model": "", "model": "", "friendly": "", "services": set(), "ports": "", "os": "",
             "found_by": []}
        for sname, data in sources.items():
            info = data.get(ip)
            if not info:
                continue
            if sname not in ("DNS", "NetBIOS"):
                d["found_by"].append(sname)
            if sname == "DNS":
                d["dns"] = d["dns"] or info
            elif sname == "NetBIOS":
                d["netbios"] = info
            else:
                for k, v in info.items():
                    if k == "services":
                        d["services"] |= v
                    elif v and not d.get(k):
                        d[k] = v
        if ip == self_ip:
            d["mac"] = d["mac"] or self_mac
            d["found_by"].append("this host")
        d.update(detail.get(ip, {}))
        d["type"] = guess_type(d)
        devices[ip] = d

    # ---- history, keyed by MAC (IP for the rare device without one)
    out = os.path.abspath(a.out)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    json_path = os.path.splitext(out)[0] + ".json"
    hist = load_previous(json_path, out)
    now = NOW.isoformat(sep=" ")
    # One row per MAC, so DHCP address changes don't create duplicates - except
    # when one MAC answers for several IPs (Wi-Fi extenders / mesh nodes do this
    # for their clients): then keep each IP as its own row.
    mac_count = Counter(d["mac"] for d in devices.values() if d["mac"])
    for d in devices.values():
        shared = d["mac"] and mac_count[d["mac"]] > 1
        key = (f"{d['mac']}|{d['ip']}" if shared else d["mac"]) or f"ip:{d['ip']}"
        if shared:
            d["found_by"].append("shared MAC (extender/proxy?)")
        h = hist.get(key, {"first_seen": now, "friendly_user": "", "notes": ""})
        h.update({
            "ip": d["ip"], "mac": d["mac"], "vendor": d["vendor"], "type": d["type"],
            "hostname": d["mdns"].removesuffix(".local") if d["mdns"] else (d["dns"] or d["netbios"] or d["friendly"] or d["upnp_name"]),
            "mdns": d["mdns"], "dns": d["dns"], "netbios": d["netbios"],
            "upnp": " - ".join(x for x in (d["friendly"] or d["upnp_name"], d["model"] or d["upnp_model"]) if x),
            "services": ", ".join(sorted(d["services"])), "ports": d["ports"], "os": d["os"],
            "private_mac": "yes" if is_private_mac(d["mac"]) else "", "found_by": ", ".join(d["found_by"]),
            "last_seen": now, "online": True,
        })
        hist[key] = h
    cutoff = NOW - dt.timedelta(days=a.keep_days)
    for key in list(hist):
        if hist[key].get("last_seen", now) != now:
            hist[key]["online"] = False
            if dt.datetime.fromisoformat(hist[key]["last_seen"]) < cutoff:
                del hist[key]
    with open(json_path, "w") as f:
        json.dump(hist, f, indent=1, sort_keys=True)

    write_xlsx(out, hist, a, self_ip, iface)
    online = sum(1 for h in hist.values() if h["online"])
    log(f"\n{online} devices online, {len(hist) - online} remembered offline -> {out}")
    print(json.dumps({"report": out, "online": online, "offline": len(hist) - online,
                      "by_source": {k: len(v) for k, v in sources.items()}}))


def write_xlsx(path, hist, a, self_ip, iface):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Devices"
    head_fill = PatternFill("solid", fgColor="1F3864")
    yours_fill = PatternFill("solid", fgColor="FFF2CC")
    off_font = Font(color="808080", italic=True)
    ws.append([c for c, _ in COLUMNS])
    for i, (_, w) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30

    rows = sorted(hist.values(), key=lambda h: (not h["online"], ip_key(h["ip"])))
    for h in rows:
        ws.append([
            "online" if h["online"] else "offline", h["ip"], h["hostname"], h["mac"], h["vendor"], h["type"],
            h.get("friendly_user", ""), h.get("notes", ""), h["mdns"], h["dns"], h["netbios"], h["upnp"],
            h["services"], h["ports"], h["os"], h["private_mac"], h["found_by"], h["first_seen"], h["last_seen"],
        ])
        r = ws.max_row
        for col in (7, 8):
            ws.cell(row=r, column=col).fill = yours_fill
        if not h["online"]:
            for col in range(1, len(COLUMNS) + 1):
                ws.cell(row=r, column=col).font = off_font
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    # ---- summary sheet
    s = wb.create_sheet("Summary")
    online = [h for h in hist.values() if h["online"]]
    s.append(["Network inventory"]); s["A1"].font = Font(bold=True, size=14)
    for k, v in [("Scanned", NOW.isoformat(sep=" ")), ("Subnet", a.subnet), ("Scanned from", f"{self_ip} ({iface})"),
                 ("DNS server", a.dns or "system default"), ("Devices online", len(online)),
                 ("Remembered offline", len(hist) - len(online)),
                 ("Private (random) MACs", sum(1 for h in online if h["private_mac"]))]:
        s.append([k, v])
    for title, counter in (("By device type", Counter(h["type"] or "Unknown" for h in online)),
                           ("By vendor", Counter(h["vendor"] or "Unknown" for h in online))):
        s.append([])
        s.append([title, "Count"])
        s.cell(row=s.max_row, column=1).font = Font(bold=True)
        s.cell(row=s.max_row, column=2).font = Font(bold=True)
        for name, n in counter.most_common():
            s.append([name, n])
    s.column_dimensions["A"].width = 42
    s.column_dimensions["B"].width = 28
    wb.save(path)


if __name__ == "__main__":
    main()
